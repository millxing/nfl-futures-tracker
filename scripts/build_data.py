#!/usr/bin/env python3
"""Convert snapshot xlsx files into docs/data/data.json for the site.

Usage: python3 scripts/build_data.py
Reads every snapshots/nfl_markets_snapshot_*.xlsx (Polymarket sheet only),
sorted chronologically by the datetime in the filename.
Optional snapshots/labels.json maps snapshot id -> friendly label,
e.g. {"2026-09-10_2000": "Pre Week 1"}.
"""

import glob
import json
import os
import re
import sys
from datetime import datetime, timezone

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SNAPSHOT_DIR = os.path.join(ROOT, "snapshots")
OUT_PATH = os.path.join(ROOT, "docs", "data", "data.json")
SHEET = "Polymarket series"

# Display order: team series first (macro -> divisions), then individual awards.
SERIES_ORDER = [
    "Pro Football: 2027 Champion",
    "Pro Football: 2027 AFC Champion",
    "Pro Football: 2027 NFC Champion",
    "Pro Football: Team to Make Postseason",
    "Pro Football: AFC East Champion",
    "Pro Football: AFC North Champion",
    "Pro Football: AFC South Champion",
    "Pro Football: AFC West Champion",
    "Pro Football: NFC East Champion",
    "Pro Football: NFC North Champion",
    "Pro Football: NFC South Champion",
    "Pro Football: NFC West Champion",
    "Pro Football: 2026 MVP Winner",
    "Pro Football: 2026-27 AP Offensive Player of the Year Winner",
    "Pro Football: 2026-27 AP Defensive Player of the Year Winner",
    "Pro Football: 2026-27 AP Offensive Rookie of the Year Winner",
    "Pro Football: 2026-27 AP Defensive Rookie of the Year Winner",
    "Pro Football: 2026-27 AP Coach of the Year Winner",
    "Pro Football: 2026-27 AP Comeback Player of the Year Winner",
]
TEAM_SERIES = set(SERIES_ORDER[:12])

# Tracked on Polymarket but too illiquid for prices to read as probabilities.
EXCLUDED_SERIES = {
    "Pro Football: 2026-27 AFC #1 Seed",
    "Pro Football: 2026-27 NFC #1 Seed",
}

FNAME_RE = re.compile(
    r"nfl_markets_snapshot_(\d{4}-\d{2}-\d{2})_(\d{4})_E[SD]T\.xlsx$"
)


def snapshot_files():
    files = []
    for path in glob.glob(os.path.join(SNAPSHOT_DIR, "nfl_markets_snapshot_*.xlsx")):
        base = os.path.basename(path)
        if base.startswith("~$"):
            continue
        m = FNAME_RE.search(base)
        if not m:
            print(f"WARNING: skipping unrecognized filename {base}", file=sys.stderr)
            continue
        dt = datetime.strptime(f"{m.group(1)} {m.group(2)}", "%Y-%m-%d %H%M")
        files.append((dt, f"{m.group(1)}_{m.group(2)}", path))
    files.sort()
    return files


def auto_label(dt):
    hour = dt.strftime("%I:%M %p").lstrip("0")
    return f"{dt.strftime('%b %-d')}, {hour} ET"


def read_snapshot(path):
    """Return {market_id: row-dict} for the Polymarket sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # title
    next(rows)  # created
    header = [str(h) if h is not None else "" for h in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    out = {}
    for row in rows:
        if row is None or row[idx["market_id"]] is None:
            continue
        if str(row[idx["event_title"]]).strip() in EXCLUDED_SERIES:
            continue
        mid = str(row[idx["market_id"]])
        price = row[idx["price"]]
        status = row[idx["data_status"]]
        out[mid] = {
            "series": str(row[idx["event_title"]]).strip(),
            "name": str(row[idx["outcome"]]).strip(),
            "question": str(row[idx["market_question"]]).strip(),
            "price": round(float(price) * 100, 1)
            if status == "ok" and price is not None
            else None,
        }
    wb.close()
    return out


def main():
    files = snapshot_files()
    if not files:
        sys.exit("No snapshot files found in snapshots/")

    labels = {}
    labels_path = os.path.join(SNAPSHOT_DIR, "labels.json")
    if os.path.exists(labels_path):
        with open(labels_path) as f:
            labels = json.load(f)

    snapshots = []
    per_snapshot = []
    for dt, sid, path in files:
        snapshots.append(
            {
                "id": sid,
                "label": labels.get(sid, auto_label(dt)),
                "et": dt.strftime("%Y-%m-%d %H:%M"),
            }
        )
        per_snapshot.append(read_snapshot(path))
        print(f"  {sid}: {len(per_snapshot[-1])} markets")

    # Collect all markets across snapshots; metadata from latest appearance.
    meta = {}
    for snap in per_snapshot:
        for mid, row in snap.items():
            meta[mid] = row

    series_seen = {m["series"] for m in meta.values()}
    unknown = series_seen - set(SERIES_ORDER)
    if unknown:
        print(f"WARNING: series not in SERIES_ORDER (appended): {unknown}", file=sys.stderr)
    series = [s for s in SERIES_ORDER if s in series_seen] + sorted(unknown)
    series_rank = {s: i for i, s in enumerate(series)}

    markets = []
    for mid, m in meta.items():
        markets.append(
            {
                "id": mid,
                "series": m["series"],
                "name": m["name"],
                "question": m["question"],
                "prices": [snap.get(mid, {}).get("price") for snap in per_snapshot],
            }
        )
    markets.sort(key=lambda m: (series_rank[m["series"]], m["name"]))

    data = {
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "snapshots": snapshots,
        "series": series,
        "teamSeries": sorted(TEAM_SERIES & series_seen, key=series_rank.get),
        "markets": markets,
    }

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w") as f:
        json.dump(data, f, separators=(",", ":"))

    print(f"Wrote {OUT_PATH}")
    print(f"  snapshots: {len(snapshots)}")
    print(f"  series:    {len(series)}")
    print(f"  markets:   {len(markets)}")


if __name__ == "__main__":
    main()
